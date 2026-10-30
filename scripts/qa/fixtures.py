from pathlib import Path
from PIL import Image
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen.canvas import Canvas
p=Path('qa-output');p.mkdir(exist_ok=True)
for name,color in [('blue',(0,0,255,255)),('transparent',(0,0,0,0))]:Image.new('RGBA',(120,80),color).save(p/(name+'.png'))
c=Canvas(str(p/'fixture.pdf'))
for i in range(1,4):c.drawString(72,720,f'Vexa page {i} test marker');c.showPage()
c.save();c=Canvas(str(p/'blank.pdf'));c.showPage();c.save()
d=Document();d.add_paragraph('WORD-START');t=d.add_table(rows=2,cols=2);t.cell(0,0).text='Name';t.cell(0,1).text='Value';t.cell(1,0).text='Vexa';t.cell(1,1).text='42'
for i in range(80):d.add_paragraph(f'WORD-ROW-{i:03} Long document text A & B')
d.add_paragraph('WORD-END');d.save(p/'fixture.docx')
w=Workbook();s=w.active;s.title='Data';s.append(['Name','Value'])
for i in range(140):s.append([f'EXCEL-ROW-{i:03}',i])
s.append(['EXCEL-END',140]);w.create_sheet('Summary').append(['SECOND-SHEET',42]);w.save(p/'fixture.xlsx')
r=Presentation()
for n in range(2):
 sl=r.slides.add_slide(r.slide_layouts[6]);tf=sl.shapes.add_textbox(Inches(.5),Inches(.5),Inches(8),Inches(6)).text_frame;tf.text=f'SLIDE-{n}'
 for i in range(65):tf.add_paragraph().text=f'PPT-{n}-ROW-{i:03}'
 tf.add_paragraph().text=f'PPT-{n}-END'
r.save(p/'fixture.pptx')
(p/'corrupt.pdf').write_text('not a PDF')
Image.new('RGB',(120,80),'blue').save(p/'blue.jpg')
d=Document();d.add_paragraph('తెలుగు పరీక్ష — नमस्ते भारत');d.add_paragraph('Unicode end marker');d.save(p/'unicode.docx')
from pypdf import PdfWriter
from pypdf.generic import RectangleObject
w=PdfWriter()
for angle in [0,90,180,270]:
 page=w.add_blank_page(width=400,height=600);page.cropbox=RectangleObject([20,40,380,560]);page.rotate(angle)
w.write(p/'rotated-crop.pdf')
Image.new('RGB',(2000,1000),'red').save(p/'large-logo.png')
