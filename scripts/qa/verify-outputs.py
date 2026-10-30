from pathlib import Path
from pypdf import PdfReader
from PIL import Image
from docx import Document
from openpyxl import load_workbook
from zipfile import ZipFile
import json
p=Path('qa-output');out=[]
for width in [390,1440]:
 def file(prefix):
  files=[f for f in p.glob(f'{prefix}-{width}-*') if not any(f.name.endswith(x) for x in ['-0.png','-3.png','-final.png','-preview.png'])];assert files,prefix;return files[0]
 for prefix,count in [('merge',4),('split',2),('rotate',3),('compress',3),('organize-pdf',3),('repair-pdf',3),('watermark-pdf',3)]:
  r=PdfReader(file(prefix),strict=True);assert len(r.pages)==count,(prefix,len(r.pages));assert 'Vexa page' in ''.join(x.extract_text() or '' for x in r.pages);out.append({'tool':prefix,'width':width,'pages':count,'text':True})
 r=PdfReader(file('protected'));assert r.is_encrypted;assert not r.decrypt('wrong');assert r.decrypt('Vexa-password42');assert len(r.pages)==3
 for prefix,markers in [('excel-to-pdf',['EXCEL-END','SECOND-SHEET']),('ppt-to-pdf',['PPT-0-END','PPT-1-END']),('html',['VEXA HTML','Output marker']),('edit',['VEXA EDIT VERIFIED'])]:
  r=PdfReader(file(prefix),strict=True);text=''.join(x.extract_text() or '' for x in r.pages);assert all(m in text for m in markers),(prefix,text[-100:]);out.append({'tool':prefix,'width':width,'pages':len(r.pages),'endMarkers':True})
 r=PdfReader(file('word-to-pdf'));assert len(r.pages)>1;assert all(len(page.images)>0 for page in r.pages)
 r=PdfReader(file('sign'),strict=True);assert len(r.pages)==3
 r=PdfReader(file('unlock'),strict=True);assert len(r.pages)==3;assert all(len(page.images)>0 for page in r.pages)
 for prefix in ['image-to-pdf','jpg-to-pdf','png-to-pdf']:assert len(PdfReader(file(prefix)).pages)==1
 for prefix in ['pdf-to-word','excel-to-word']:
  d=Document(file(prefix));text=' '.join(x.text for x in d.paragraphs)+' '.join(c.text for t in d.tables for row in t.rows for c in row.cells);assert ('Vexa page' if prefix=='pdf-to-word' else 'EXCEL-END') in text
 for prefix in ['word-to-excel','pdf-to-excel']:
  w=load_workbook(file(prefix));text=str([list(s.values) for s in w]);assert ('Vexa' in text),prefix
 with ZipFile(file('pdf-to-powerpoint')) as z:assert len([n for n in z.namelist() if n.lower().endswith('.jpg')])==3
 assert 'Vexa page' in file('pdf-to-html').read_text()
 for prefix in ['pdf-to-jpg','pdf-to-png','pdf-to-image','image-resizer','compress-image','aadhaar-photo-resize','govt-job-photo-resize','passport-photo-resize','whatsapp-dp-resize','qr']:
  im=Image.open(file(prefix));im.load();assert im.width>0 and im.height>0
 im=Image.open(file('image-format-converter'));im.load();assert im.format=='WEBP';assert im.convert('RGBA').getpixel((0,0))[3]==0,'WebP transparency lost'
 out.append({'width':width,'officeFilesOpen':True,'slideImages':3,'imagesDecode':True,'webpTransparency':True})
p.joinpath('output-verification.json').write_text(json.dumps(out,indent=2));print('Independent output validation PASS:',len(out),'checks')
